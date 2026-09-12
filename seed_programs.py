"""One-time seed of data/programs.csv from the hand-built spreadsheet.

Run once at setup. After this, programs.csv is the source of truth and the xlsx is a
build artifact (spec section 4); re-running would overwrite tracking state, so it
refuses to clobber an existing file unless --force is passed.
"""
from __future__ import annotations

import argparse
import pathlib
import sys

import openpyxl

from core import paths
from persist import store

# xlsx header -> programs.csv column. The spreadsheet uses human-readable headers;
# the CSV uses the snake_case names from spec section 4.
HEADER_MAP = {
    "Category": "category",
    "Name": "name",
    "Website": "website",
    "Applications Open": "applications_open",
    "Target Years in College": "target_years",
    "Eligible for You": "eligible",
    "Notes": "notes",
}

# Tier 1 repos map onto existing "Tracker / resource" rows; match on the website URL
# so the monitored rows are the ones already researched by hand.
REPO_URL_TO_SOURCE_ID = {
    "github.com/northwesternfintech/2027quantinternships": "nuft-2027",
    "github.com/simplifyjobs/summer2027-internships": "simplify-2027",
    "github.com/jose-gael-cruz-lopez/underclassmen-opportunities": "underclassmen-cruz",
    "github.com/zapplyjobs/underclassmen-internships": "underclassmen-zapply",
    "github.com/luisae/opportunities": "luisae",
}


def source_id_for(website: str) -> str:
    normalised = (website or "").lower().rstrip("/").removeprefix("https://").removeprefix("http://")
    return REPO_URL_TO_SOURCE_ID.get(normalised, "")


def seed(xlsx: pathlib.Path) -> list[dict[str, str]]:
    worksheet = openpyxl.load_workbook(xlsx, data_only=True)["Programs"]
    headers = [cell.value for cell in worksheet[1]]
    missing = [h for h in HEADER_MAP if h not in headers]
    if missing:
        raise SystemExit(f"seed file is missing expected columns: {missing}")

    rows: list[dict[str, str]] = []
    for excel_row in range(2, worksheet.max_row + 1):
        row = {column: "" for column in paths.PROGRAM_COLUMNS}
        for header, column in HEADER_MAP.items():
            value = worksheet.cell(excel_row, headers.index(header) + 1).value
            row[column] = "" if value is None else str(value).strip()
        if not row["name"]:
            continue
        row["source_id"] = source_id_for(row["website"])
        row["status"] = "unknown"
        row["muted"] = "false"
        # last_checked / last_changed / snapshot_hash stay blank until a real check
        # runs, and first_seen stays blank because these rows are hand-researched,
        # not auto-discovered (spec section 4).
        rows.append(row)
    return rows


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--xlsx",
        type=pathlib.Path,
        default=paths.ROOT.parent / "quant_math_cs_programs_freshman.xlsx",
        help="path to the seed spreadsheet",
    )
    parser.add_argument("--force", action="store_true", help="overwrite an existing programs.csv")
    args = parser.parse_args()

    if paths.PROGRAMS_CSV.exists() and not args.force:
        print(f"{paths.PROGRAMS_CSV} already exists; refusing to overwrite tracking state.")
        print("Pass --force if you really mean to re-seed from the spreadsheet.")
        return 1
    if not args.xlsx.exists():
        print(f"seed spreadsheet not found: {args.xlsx}")
        return 1

    rows = seed(args.xlsx)
    store.write_programs(rows)
    monitored = sum(1 for row in rows if row["source_id"])
    print(f"wrote {len(rows)} rows to {paths.PROGRAMS_CSV} ({monitored} wired to a source)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
