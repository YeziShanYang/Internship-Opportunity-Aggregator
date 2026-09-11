"""Regenerate the two workbooks in out/ from the CSVs (spec section 4).

Both are build artifacts, never hand-edited. Git treats a spreadsheet as an opaque zip
of XML, so committing one as the source of truth would make every run look like the
whole file changed and destroy the audit trail. The CSVs are canonical; this script
renders them, preserving the seed file's formatting: Arial, a colour-coded eligibility
column, autofilter, frozen panes and a Legend sheet.

**Two workbooks, split by audience.** This split was asked for directly: programs.xlsx
had grown into a dump of all 187 tracked programmes, which made it useless as the thing
the owner actually opens to decide what to go and check. Two categories of row are pure
bloat *for that purpose*, even though both are worth keeping as data:

  * things he cannot apply to (`eligible` NO or STALE), and
  * things the tracker already watches for him, which will arrive in the daily digest
    whether or not he ever opens a spreadsheet.

So:

  out/programs.xlsx   HIS file. Only what he must chase himself.
                      Check By Hand / Manual Watch / Priority / Legend / Left Out
  out/tracked.xlsx    the TOOL's bookkeeping, moved out of his way.
                      All Programs / Sources / Applied / Discovered

The Left Out sheet is not decoration. This project's standing rule is that every filter
reports what it removed, because a filter that hides silently is how a source goes blind
without anyone noticing -- it has happened twice here. Left Out names every excluded
programme and the exact reason, so the narrowing of his file is auditable rather than
a matter of trust.
"""
from __future__ import annotations

import csv
import pathlib
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import state

ARIAL = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)
BODY_FONT = Font(name=ARIAL, size=10)
BODY_ALIGN = Alignment(vertical="top", wrap_text=True)
LINK_FONT = Font(name=ARIAL, size=10, color="0563C1", underline="single")

# Exactly the palette from the seed spreadsheet.
GREEN = (PatternFill("solid", fgColor="C6EFCE"), "006100")
RED = (PatternFill("solid", fgColor="FFC7CE"), "9C0006")
AMBER = (PatternFill("solid", fgColor="FFEB9C"), "9C6500")
BLUE = (PatternFill("solid", fgColor="DDEBF7"), "1F4E79")

ELIGIBLE_STYLES = {
    "YES": GREEN,
    "USE": GREEN,
    "USE WEEKLY": GREEN,
    "NO": RED,
    "STALE": RED,
    "LATER": AMBER,
    "LOW VALUE": AMBER,
    "APPLIED - too early": AMBER,
    "CHECK": BLUE,
    "INVITE": BLUE,
    "VIA CLUB": BLUE,
}

# Cannot be applied to at all, now or later. These leave his file entirely.
CANNOT_APPLY = {"NO", "STALE"}

# Sort order for the hand-check sheet: what to do something about first. Everything
# outside this list sorts last but is still present.
ELIGIBLE_ORDER = [
    "YES", "CHECK", "INVITE", "VIA CLUB", "APPLIED - too early",
    "LATER", "LOW VALUE", "USE", "USE WEEKLY",
]

# Methods that watch one specific page, so a URL match against them is real evidence
# that a programme is covered. Slug-based methods (greenhouse/lever/ashby) are absent
# on purpose -- see `covered_by`.
PAGE_METHODS = ("page_text", "workday", "phenom")

HAND_CHECK_HEADERS = [
    ("eligible", "Eligible for You", 17),
    ("name", "Name", 46),
    ("category", "Category", 20),
    ("applications_open", "When it opens", 30),
    ("target_years", "Target Years in College", 26),
    ("website", "Website", 56),
    ("notes", "Notes", 90),
]

LEFT_OUT_HEADERS = [
    ("why", "Why it is not on your list", 40),
    ("name", "Name", 46),
    ("category", "Category", 20),
    ("eligible", "Eligible for You", 17),
    ("website", "Website", 52),
    ("notes", "Notes", 78),
]

PROGRAM_HEADERS = [
    ("category", "Category", 20),
    ("name", "Name", 44),
    ("website", "Website", 56),
    ("applications_open", "Applications Open", 30),
    ("target_years", "Target Years in College", 30),
    ("eligible", "Eligible for You", 17),
    ("notes", "Notes", 78),
    ("source_id", "Source ID", 22),
    ("status", "Status", 12),
    ("last_checked", "Last Checked", 22),
    ("last_changed", "Last Changed", 22),
    ("first_seen", "First Seen", 14),
    ("muted", "Muted", 8),
    ("snapshot_hash", "Snapshot Hash", 18),
]

MANUAL_HEADERS = [
    ("name", "Name", 44),
    ("category", "Category", 20),
    ("coverage", "Will this tool cover it?", 26),
    ("why_manual", "Why it is manual", 76),
    ("how_to_check", "How to check", 54),
    ("when_to_check", "When to check", 46),
    ("url", "URL", 56),
]

PRIORITY_HEADERS = [
    ("rank", "#", 5),
    ("band", "Band", 34),
    ("name", "Name", 46),
    ("category", "Category", 20),
    ("value", "Value", 13),
    ("probability", "Odds", 13),
    ("watch_window", "Watch window", 38),
    ("will_this_tool_alert_you", "Will this tool alert you?", 30),
    ("why", "Why it earns the slot", 96),
    ("url", "URL", 52),
]

SOURCE_HEADERS = [
    ("source_id", "Source ID", 34),
    ("method", "Method", 15),
    ("signal", "Signal", 10),
    ("program_names", "Programs", 34),
    ("url", "URL or slug", 60),
    ("last_success", "Last Success", 22),
    ("consecutive_failures", "Fails", 7),
    ("notes", "Notes", 100),
]

APPLIED_HEADERS = [
    ("title", "What you ticked off", 90),
    ("marked_at", "Marked applied", 18),
    ("key", "Key", 14),
]

DISCOVERED_HEADERS = [
    ("status", "Status", 12),
    ("title", "Title", 40),
    ("kind", "Kind", 10),
    ("key", "Key", 28),
    ("url", "URL or slug", 50),
    ("evidence", "Evidence / why", 100),
    ("first_proposed", "First proposed", 16),
    ("last_proposed", "Last proposed", 16),
]

LEGEND_ROWS = [
    ("THIS FILE", "What it is"),
    (
        "out/programs.xlsx",
        "Yours. Every row here is something YOU have to go and look at - the tracker "
        "either cannot see it, or cannot tell whether it applies to you. If a row is "
        "here, nobody else is watching it.",
    ),
    (
        "out/tracked.xlsx",
        "The tool's own records: the full programme list, all watched sources and their "
        "health, what you have ticked off, and what weekly discovery has proposed or "
        "permanently rejected. You never need to open it to decide what to apply to.",
    ),
    (
        "What was removed",
        "Two things, both listed with reasons on the Left Out sheet: programmes you "
        "cannot apply to (NO / STALE), and programmes a watched source already covers - "
        "those reach you in the daily digest instead.",
    ),
    (None, None),
    ("Sheet", "What it is for"),
    (
        "Check By Hand",
        "The list. Sorted so the actionable codes come first: YES, then CHECK, then "
        "invite-only, then things for a later year.",
    ),
    (
        "Manual Watch",
        "Opportunities this tool provably cannot alert you about - sites that block "
        "automation, competitions announced on Instagram or a listserv before the "
        "website changes, and pages with nothing to diff. Each row says how and when to "
        "check by hand.",
    ),
    (
        "Priority",
        "A ranked shortlist to keep an eye on yourself in case the tracker breaks or goes "
        "quiet. Band A is high value with real selection risk; Band B is "
        "Stanford-internal, where your odds are genuinely best; Band C is open entry, "
        "where showing up is the only gate.",
    ),
    (
        "Left Out",
        "The audit trail for this file's filtering. Nothing is hidden from you silently; "
        "if you disagree with an exclusion, the reason is written down next to it.",
    ),
    (None, None),
    ("Column", "What it means"),
    (
        "Eligible for You",
        "Filtered for your stated profile: first-year (class of 2030), male, not from a "
        "marginalized group, US-based, at Stanford.",
    ),
    (None, None),
    ("Code", "Meaning"),
    ("YES", "Verified open to you now. Apply."),
    ("LATER", "Right program, wrong year (or wrong identity-neutral stage). Diary it."),
    (
        "CHECK",
        "Program exists but eligibility is not published. Open the page yourself before "
        "assuming.",
    ),
    ("INVITE", "Entry is by invitation via campus recruiting, not open application."),
    ("VIA CLUB", "Entry runs through a university-sponsored team (Traders @ Stanford)."),
    ("LOW VALUE", "Open to you, but weak signal relative to the time it costs."),
    ("USE / USE WEEKLY", "A tracker or resource to monitor, not an application."),
    ("APPLIED - too early", "You have already applied and it was not yet the right cycle."),
    (
        "NO / STALE",
        "Ruled out, or the program no longer exists. These are NOT in this file - see "
        "the Left Out sheet, or All Programs in tracked.xlsx.",
    ),
    (None, None),
    (
        "Dates",
        "Dates are the PRIOR cycle's unless a 2026/2027 date is explicitly stated. Treat "
        "them as the window to start watching, not a deadline to trust.",
    ),
    (
        "Rolling firms",
        "Jane Street, NVIDIA, D. E. Shaw and Point72 review on a rolling basis and close "
        "when full. Week-one applications face the most open seats.",
    ),
    (None, None),
    (
        "Eligibility warning",
        "A programme's NAME is not its eligibility gate. Susquehanna's US Discovery "
        "Programs are described everywhere as first- and second-year, and every posting "
        "requires graduation by winter 2028 / spring 2029 - a sophomore on a four-year "
        "US degree. Read the stated graduation window, not the label.",
    ),
    (
        "Open to a first-year right now",
        "Voloridge Ascend (first and second year, 20 places) | HRT Inside HRT (first and "
        "second year) | Jane Street INSIGHT (first year) | Optiver FutureFocus (label "
        "says first and second year; page states no graduation window, so verify)",
    ),
    (None, None),
    (
        "THIS FILE IS GENERATED",
        "Edit data/programs.csv, not this spreadsheet. Anything you type here is "
        "overwritten on the next run.",
    ),
]


def _write_sheet(worksheet, headers, rows, hyperlink_columns=()) -> None:
    for index, (_, title, width) in enumerate(headers, start=1):
        cell = worksheet.cell(1, index, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.row_dimensions[1].height = 24

    for row_index, row in enumerate(rows, start=2):
        for column_index, (key, _, _) in enumerate(headers, start=1):
            value = row.get(key, "")
            cell = worksheet.cell(row_index, column_index, value)
            cell.font = BODY_FONT
            cell.alignment = BODY_ALIGN
            if key in hyperlink_columns and isinstance(value, str) and value.startswith("http"):
                cell.hyperlink = value
                cell.font = LINK_FONT
            if key == "eligible":
                style = ELIGIBLE_STYLES.get(str(value).strip())
                if style:
                    fill, colour = style
                    cell.fill = fill
                    cell.font = Font(name=ARIAL, size=10, bold=True, color=colour)

    last_column = get_column_letter(len(headers))
    last_row = len(rows) + 1
    worksheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    worksheet.freeze_panes = "B2"


def _read(path: pathlib.Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _canon(url: str) -> str:
    text = (url or "").strip().lower()
    text = re.sub(r"^https?://", "", text)
    text = re.sub(r"^www\.", "", text)
    return text.rstrip("/").split("?")[0]


def covered_by(program: dict[str, str], sources: dict[str, dict[str, str]]) -> str:
    """Which watched source already covers this programme, or "" if none does.

    Only *specific* evidence counts, and that restriction is the whole point. A first
    version matched on the registrable domain and was wrong in a way that mattered: it
    called "Jane Street Puzzles (monthly)" and "Jane Street SF Puzzle City" covered
    because janestreet.com appears in the watchlist, when the watchers are pointed at
    FTTP, INSIGHT and the programmes index. Those puzzle pages are exactly the kind of
    thing the owner has to check himself, and folding them away would have hidden them
    from the only file that was going to tell him.

    So the bar is an explicit `source_id`, or a watched page whose URL is the programme's
    URL (or a parent of it). Wrongly keeping a row costs a glance; wrongly dropping one
    costs an opportunity, which is the error this project always spends a little noise
    to avoid.

    A `source_id` pointing at a row that no longer exists returns a BROKEN marker. It
    deliberately does not count as coverage -- a dangling reference must not quietly
    remove a programme from the owner's list.
    """
    for source_id in [s.strip() for s in (program.get("source_id") or "").split(";") if s.strip()]:
        source = sources.get(source_id)
        if source is None:
            return f"BROKEN source_id={source_id}"
        if source["method"] != state.UNWATCHED_METHOD:
            return source_id
    target = _canon(program.get("website"))
    if target:
        for source in sources.values():
            if source["method"] not in PAGE_METHODS:
                continue
            page = _canon(source["url"])
            if page and (target == page or target.startswith(page + "/")):
                return source["source_id"]
    return ""


def partition(
    programs: list[dict[str, str]], sources: dict[str, dict[str, str]]
) -> tuple[list[dict[str, str]], list[dict[str, str]], list[str]]:
    """Split into (check-by-hand, left-out-with-reasons, warnings)."""
    keep: list[dict[str, str]] = []
    left_out: list[dict[str, str]] = []
    warnings: list[str] = []

    for program in programs:
        eligible = (program.get("eligible") or "").strip().upper()
        coverage = covered_by(program, sources)
        if coverage.startswith("BROKEN"):
            warnings.append(
                f"{program.get('name', '?')}: {coverage} is not a row in sources.csv. "
                "Kept on your list rather than dropped, but the reference should be fixed."
            )
            coverage = ""

        if eligible in CANNOT_APPLY:
            row = dict(program)
            row["why"] = f"You cannot apply ({eligible})"
            left_out.append(row)
        elif coverage:
            row = dict(program)
            row["why"] = f"Already watched by {coverage}"
            left_out.append(row)
        else:
            keep.append(program)

    def sort_key(program: dict[str, str]):
        eligible = (program.get("eligible") or "").strip()
        rank = ELIGIBLE_ORDER.index(eligible) if eligible in ELIGIBLE_ORDER else len(ELIGIBLE_ORDER)
        return (rank, program.get("category", ""), program.get("name", ""))

    keep.sort(key=sort_key)
    left_out.sort(key=lambda r: (r["why"], r.get("name", "")))
    return keep, left_out, warnings


def _write_legend(workbook) -> None:
    legend = workbook.create_sheet("Legend")
    legend.column_dimensions["A"].width = 26
    legend.column_dimensions["B"].width = 110
    for row_index, (left, right) in enumerate(LEGEND_ROWS, start=1):
        left_cell = legend.cell(row_index, 1, left)
        right_cell = legend.cell(row_index, 2, right)
        is_header = right in ("What it is", "What it is for", "What it means", "Meaning") or (
            left is not None and left.isupper() and left.startswith("THIS FILE IS")
        )
        for cell in (left_cell, right_cell):
            cell.font = HEADER_FONT if is_header else BODY_FONT
            cell.alignment = HEADER_ALIGN if is_header else BODY_ALIGN
            if is_header:
                cell.fill = HEADER_FILL
        if left is not None and not is_header:
            left_cell.font = Font(name=ARIAL, size=10, bold=True)


def build_owner_workbook(
    keep: list[dict[str, str]], left_out: list[dict[str, str]]
) -> pathlib.Path:
    """out/programs.xlsx -- only what the owner has to chase himself."""
    workbook = Workbook()
    _write_sheet(
        workbook.active, HAND_CHECK_HEADERS, keep, hyperlink_columns={"website"}
    )
    workbook.active.title = "Check By Hand"

    manual = _read(state.DATA / "manual.csv")
    if manual:
        _write_sheet(
            workbook.create_sheet("Manual Watch"), MANUAL_HEADERS, manual,
            hyperlink_columns={"url"},
        )

    priority = _read(state.DATA / "priority.csv")
    if priority:
        for row in priority:
            row["rank"] = (
                int(row["rank"]) if str(row.get("rank", "")).isdigit() else row.get("rank", "")
            )
        _write_sheet(
            workbook.create_sheet("Priority"), PRIORITY_HEADERS, priority,
            hyperlink_columns={"url"},
        )

    _write_legend(workbook)
    _write_sheet(
        workbook.create_sheet("Left Out"), LEFT_OUT_HEADERS, left_out,
        hyperlink_columns={"website"},
    )

    state.OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(state.OUT_XLSX)
    return state.OUT_XLSX


def build_tracked_workbook(
    programs: list[dict[str, str]], sources: list[dict[str, str]]
) -> pathlib.Path:
    """out/tracked.xlsx -- the tool's bookkeeping, kept out of the owner's way."""
    workbook = Workbook()
    _write_sheet(workbook.active, PROGRAM_HEADERS, programs, hyperlink_columns={"website"})
    workbook.active.title = "All Programs"

    _write_sheet(
        workbook.create_sheet("Sources"), SOURCE_HEADERS,
        sorted(sources, key=lambda s: (s["method"], s["source_id"])),
        hyperlink_columns={"url"},
    )

    applied = state.read_applied()
    titles = {}
    if state.APPLIED_TSV.exists():
        for line in state.APPLIED_TSV.read_text(encoding="utf-8").splitlines():
            if not line.strip() or line.startswith("#"):
                continue
            parts = line.split("\t")
            if len(parts) >= 3:
                titles[parts[0]] = parts[2]
    _write_sheet(
        workbook.create_sheet("Applied"), APPLIED_HEADERS,
        [
            {"key": key, "marked_at": marked, "title": titles.get(key, "")}
            for key, marked in sorted(applied.items(), key=lambda kv: kv[1], reverse=True)
        ],
    )

    _write_sheet(
        workbook.create_sheet("Discovered"), DISCOVERED_HEADERS,
        sorted(state.read_discovered(), key=lambda r: (r.get("status", ""), r.get("key", ""))),
        hyperlink_columns={"url"},
    )

    state.OUT_TRACKED_XLSX.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(state.OUT_TRACKED_XLSX)
    return state.OUT_TRACKED_XLSX


def build() -> list[pathlib.Path]:
    programs = state.read_programs()
    if not programs:
        raise SystemExit("data/programs.csv is empty; run seed_programs.py first")
    source_rows = _read(state.SOURCES_CSV)
    sources = {s["source_id"]: s for s in source_rows}

    keep, left_out, warnings = partition(programs, sources)
    for warning in warnings:
        print(f"warning: {warning}", file=sys.stderr)

    owner = build_owner_workbook(keep, left_out)
    tracked = build_tracked_workbook(programs, source_rows)
    print(
        f"{len(programs)} programs -> {len(keep)} to check by hand, "
        f"{len(left_out)} left out ("
        + ", ".join(
            f"{sum(1 for r in left_out if r['why'].startswith(prefix))} {label}"
            for prefix, label in (("You cannot apply", "not eligible"), ("Already watched", "covered"))
        )
        + ")"
    )
    return [owner, tracked]


def main() -> int:
    for path in build():
        print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
